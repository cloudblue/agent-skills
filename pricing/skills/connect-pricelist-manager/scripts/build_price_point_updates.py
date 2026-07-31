#!/usr/bin/env python3
"""Diff a vendor rate card against Connect price points.

Emits the per-point updates that are actually needed, plus a report of
everything that could not or should not be written. The MCP tool
`pricing_update_price_point` works one point at a time, so the number of
rows this script keeps *is* the number of calls the agent will make —
dropping unchanged rows is the whole point.

Usage:
    python build_price_point_updates.py \\
        --rate-card rate-card.csv \\
        --mapping mapping.json \\
        --points points.json \\
        --output updates.json \\
        [--report report.txt]

`points.json` is a list of price points dumped from the pricing
price-point list/get tools for the **draft** version you are editing. Their
keys are the authoritative attribute names — this script never invents one.

`mapping.json` describes the rate card's columns:

    {
      "sku_column": "Part Number",
      "currency_column": "Currency",
      "currency": "USD",
      "prices": {"Partner Cost": "cost", "MSRP": "sale_price"},
      "verify": {"Unit": "unit", "Billing Period": "period"},
      "row_filter": {"Tier": "T1", "Min Quantity": "1"},
      "point_id_key": "id",
      "point_mpn_key": "mpn"
    }

- `prices` maps a rate-card column to a **price point attribute name** you
  read off a real point. Both sides are required; nothing is guessed.
- `verify` (optional) maps a rate-card column to a point key that must
  match. Mismatches are reported, never written.
- `row_filter` (optional) keeps only rows whose columns equal these values.
  Use it to collapse a tiered card to one row per SKU *after* the human has
  chosen the band/tier rule.
- `point_id_key` / `point_mpn_key` default to `id` and `mpn`.

Output rows carry the old and new value per attribute so the agent can show
a diff before mutating anything:

    [{"point_id": "...", "mpn": "ACME-PRO-M",
      "changes": {"cost": {"from": "12.00", "to": "13.80", "pct": 15.0}}}]

Dependencies: stdlib. openpyxl only if the rate card is an XLSX
(`pip install openpyxl`).
"""
import argparse
import csv
import json
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

# Above LOOP_WARN updates, the agent should get the user's go-ahead before
# starting the loop; above LOOP_STOP it should not loop at all and should
# offer the bulk import path instead (see SKILL.md "Batch size").
LOOP_WARN = 50
LOOP_STOP = 200

# A change this large is almost always a units/precision mistake, not a
# price decision. Reported, never silently emitted as safe.
SUSPICIOUS_FACTOR = Decimal(10)

# `,` means thousands in `1,299.00` and decimals in `1.299,00`. Only these two
# unambiguous shapes are parsed; anything else with a `,` is not a number we
# are willing to guess at (a wrong guess is off by ×100 or ×1000).
THOUSANDS_COMMA = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
DECIMAL_COMMA = re.compile(r"^-?\d{1,3}(\.\d{3})*,\d+$")


def read_rate_card(path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.stderr.write(
            "openpyxl is required to read XLSX rate cards. "
            "Install with: pip install openpyxl\n",
        )
        sys.exit(2)

    # read_only: a vendor sheet can be 100k+ rows and this is one forward-only
    # pass, so there is no reason to build the full DOM.
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    try:
        headers = ["" if h is None else str(h).strip() for h in next(rows)]
        return [
            {h: ("" if v is None else v) for h, v in zip(headers, values)}
            for values in rows
            if any(v is not None and str(v).strip() for v in values)
        ]
    except StopIteration:
        return []
    finally:
        workbook.close()


def to_decimal(value):
    """Parse a price cell. Returns None for anything not a clean number."""
    if value is None:
        return None
    text = str(value).strip()
    for junk in ("$", "€", "£", " ", " "):
        text = text.replace(junk, "")
    if not text:
        return None
    if THOUSANDS_COMMA.match(text):
        text = text.replace(",", "")
    elif DECIMAL_COMMA.match(text):
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def index_points(points, id_key, mpn_key):
    by_mpn = {}
    problems = []
    for point in points:
        mpn = point.get(mpn_key)
        if not mpn:
            problems.append(f"point {point.get(id_key, '?')} has no {mpn_key!r}")
            continue
        by_mpn.setdefault(str(mpn), []).append(point)
    duplicates = [mpn for mpn, group in by_mpn.items() if len(group) > 1]
    for mpn in duplicates:
        problems.append(
            f"{mpn}: {len(by_mpn[mpn])} points share this MPN — resolve by "
            f"point id, not by MPN",
        )
    return by_mpn, problems


def keep(row, row_filter):
    return all(
        str(row.get(column, "")).strip() == str(expected).strip()
        for column, expected in row_filter.items()
    )


def build(rows, points, mapping):
    sku_column = mapping["sku_column"]
    price_columns = mapping["prices"]
    verify_columns = mapping.get("verify", {})
    row_filter = mapping.get("row_filter", {})
    id_key = mapping.get("point_id_key", "id")
    mpn_key = mapping.get("point_mpn_key", "mpn")
    currency_column = mapping.get("currency_column")
    expected_currency = mapping.get("currency")

    by_mpn, report = index_points(points, id_key, mpn_key)

    updates = []
    unchanged = 0
    seen_skus = set()

    for line, row in enumerate(rows, start=2):
        if row_filter and not keep(row, row_filter):
            continue

        sku = str(row.get(sku_column, "")).strip()
        if not sku:
            report.append(f"row {line}: no value in {sku_column!r} — skipped")
            continue
        if sku in seen_skus:
            report.append(
                f"row {line}: duplicate row for {sku} — only its FIRST row was "
                f"used. Set row_filter to choose the tier/band explicitly.",
            )
            continue

        # Before the dedup bookkeeping: a multi-currency card has one row per
        # SKU per currency, and a skipped wrong-currency row must not make the
        # right-currency row look like a duplicate.
        if currency_column and expected_currency:
            currency = str(row.get(currency_column, "")).strip()
            if currency and currency.upper() != expected_currency.upper():
                report.append(
                    f"row {line}: {sku} is priced in {currency}, price list is "
                    f"{expected_currency} — skipped, no conversion",
                )
                continue
        seen_skus.add(sku)

        group = by_mpn.get(sku)
        if not group:
            report.append(
                f"row {line}: {sku} matches no price point — item is not in "
                f"this price list (catalog gap, not a pricing fix)",
            )
            continue
        if len(group) > 1:
            report.append(
                f"row {line}: {sku} matches {len(group)} price points — resolve "
                f"by point id; no update emitted for an ambiguous MPN",
            )
            continue
        point = group[0]

        for column, point_key in verify_columns.items():
            source = str(row.get(column, "")).strip()
            target = str(point.get(point_key, "")).strip()
            if source and target and source.lower() != target.lower():
                report.append(
                    f"row {line}: {sku} {column}={source!r} but point "
                    f"{point_key}={target!r} — check the match before writing",
                )

        changes = {}
        for column, attribute in price_columns.items():
            new = to_decimal(row.get(column))
            if new is None:
                report.append(
                    f"row {line}: {sku} {column}={row.get(column)!r} is not a "
                    f"number — skipped, not coerced",
                )
                continue
            if new < 0:
                report.append(
                    f"row {line}: {sku} {column} is negative ({new}) — not a "
                    f"Connect price, escalate",
                )
                continue
            old = to_decimal(point.get(attribute))
            if old is not None and old == new:
                continue
            entry = {"from": None if old is None else str(old), "to": str(new)}
            # The least verifiable rows are the ones with no usable old value —
            # usually a wrong attribute name in mapping.json — so they get a
            # report line of their own rather than sliding through silently.
            if old is None:
                report.append(
                    f"row {line}: {sku} point has no parseable {attribute!r} "
                    f"({point.get(attribute)!r}) — check mapping.json before "
                    f"writing {new}",
                )
            elif old > 0:
                entry["pct"] = float(
                    ((new - old) / old * 100).quantize(Decimal("0.01")),
                )
                if new > 0 and (
                    new > old * SUSPICIOUS_FACTOR or old > new * SUSPICIOUS_FACTOR
                ):
                    report.append(
                        f"row {line}: {sku} {attribute} {old} -> {new} is an "
                        f"order-of-magnitude change — confirm before applying",
                    )
            if new == 0 and old is not None and old != 0:
                report.append(
                    f"row {line}: {sku} {attribute} drops to 0 — confirm "
                    f"this is intended",
                )
            elif old == 0 and new > 0:
                report.append(
                    f"row {line}: {sku} {attribute} was 0, now {new} — a point "
                    f"priced from zero, confirm this is intended",
                )
            changes[attribute] = entry

        if not changes:
            unchanged += 1
            continue

        updates.append(
            {"point_id": point.get(id_key), "mpn": sku, "changes": changes},
        )

    return updates, unchanged, report


def summarize(updates, unchanged, report):
    lines = [
        f"updates to apply:  {len(updates)}",
        f"unchanged, skipped: {unchanged}",
        f"needs attention:    {len(report)}",
    ]
    if len(updates) > LOOP_STOP:
        lines.append(
            f"WARNING: {len(updates)} per-point calls is past the {LOOP_STOP} "
            f"threshold. Do not loop. Narrow the scope, use one adjustment if "
            f"the change is uniform, or use the Connect UI bulk price import.",
        )
    elif len(updates) > LOOP_WARN:
        lines.append(
            f"NOTE: {len(updates)} per-point calls. Tell the user the call "
            f"count and get a go-ahead before starting the loop.",
        )
    if report:
        lines.append("")
        lines.extend(report)
    return "\n".join(lines)


def self_test():
    """Cheap asserts over the parsing and skip rules. `--self-test`."""
    assert to_decimal("$1,299.00") == Decimal("1299.00")
    assert to_decimal("€ 1.299,00") == Decimal("1299.00")
    assert to_decimal("12,50") == Decimal("12.50")
    assert to_decimal("12.50") == Decimal("12.50")
    assert to_decimal("1,23,45") is None, "ambiguous commas must not be coerced"
    assert to_decimal("n/a") is None

    mapping = {
        "sku_column": "SKU",
        "currency_column": "Currency",
        "currency": "USD",
        "prices": {"Cost": "cost"},
    }
    points = [{"id": "PP-1", "mpn": "A", "cost": "10.00"}]
    rows = [
        {"SKU": "A", "Currency": "EUR", "Cost": "9.00"},
        {"SKU": "A", "Currency": "USD", "Cost": "11.00"},
    ]
    updates, _, _ = build(rows, points, mapping)
    assert [u["point_id"] for u in updates] == ["PP-1"], updates
    assert updates[0]["changes"]["cost"]["to"] == "11.00", updates

    dupes = [
        {"id": "PP-1", "mpn": "A", "cost": "10.00"},
        {"id": "PP-2", "mpn": "A", "cost": "10.00"},
    ]
    updates, _, report = build(rows[1:], dupes, mapping)
    assert updates == [], updates
    assert any("resolve by point id" in line for line in report), report

    # No usable old value is the least verifiable case: it must be reported.
    for old_value, expected in (
        (None, "check mapping.json"),
        ("", "check mapping.json"),
        ("0", "was 0, now"),
    ):
        _, _, report = build(rows[1:], [{"id": "PP-1", "mpn": "A", "cost": old_value}], mapping)
        assert any(expected in line for line in report), (old_value, report)

    # Dropping a real price to 0 reports once, not twice.
    zeroed = [{"SKU": "A", "Currency": "USD", "Cost": "0"}]
    _, _, report = build(zeroed, points, mapping)
    assert [line for line in report if "drops to 0" in line], report
    assert not [line for line in report if "order-of-magnitude" in line], report
    print("self-test: ok")


def main():
    if "--self-test" in sys.argv:
        self_test()
        return
    parser = argparse.ArgumentParser(
        description="Diff a vendor rate card against Connect price points.",
    )
    parser.add_argument("--rate-card", required=True, type=Path)
    parser.add_argument(
        "--mapping",
        required=True,
        type=Path,
        help="JSON column mapping; see this script's docstring.",
    )
    parser.add_argument(
        "--points",
        required=True,
        type=Path,
        help="JSON list of price points from the draft version.",
    )
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument(
        "--report",
        type=Path,
        help="Optional path for the attention report (also printed).",
    )
    args = parser.parse_args()

    mapping = json.loads(args.mapping.read_text(encoding="utf-8"))
    for required in ("sku_column", "prices"):
        if required not in mapping:
            sys.stderr.write(f"mapping.json is missing {required!r}\n")
            sys.exit(1)

    points = json.loads(args.points.read_text(encoding="utf-8"))
    if isinstance(points, dict):
        points = points.get("items", [])
    if not isinstance(points, list):
        sys.stderr.write("points.json must contain a JSON list of points\n")
        sys.exit(1)

    rows = read_rate_card(args.rate_card)
    updates, unchanged, report = build(rows, points, mapping)

    args.output.write_text(
        json.dumps(updates, indent=2) + "\n",
        encoding="utf-8",
    )
    summary = summarize(updates, unchanged, report)
    if args.report:
        args.report.write_text(summary + "\n", encoding="utf-8")
    print(f"Read {len(rows)} rate-card row(s) -> {args.output}")
    print(summary)


if __name__ == "__main__":
    main()
